"""Deterministic writers for small sample documents: text, PDF, DOCX, XLSX and PPTX.

Used by ``make_demo_corpus.py`` (the committed eval corpus) and by the test
suite (a tiny corpus per test). No new dependencies: PDFs come from a minimal
hand-written writer, Office files from the parser libraries the backend
already depends on.

Reproducibility: text and PDF output is byte-identical on every run. Office
files are zip archives; their member timestamps and document properties are
pinned so a rerun on the same machine gives identical bytes, but compressed
bytes may differ between zlib builds. Compare their *parsed text*, not bytes.
"""

from __future__ import annotations

import io
import re
import zipfile
from collections.abc import Mapping, Sequence
from datetime import datetime
from pathlib import Path
from typing import Any

FIXED_TIME = datetime(2026, 1, 1, 0, 0, 0)
AUTHOR = "Fernhill Outdoor Co. (fictional)"
FOOTER = "Fictional sample data"

_ZIP_TIME = (2026, 1, 1, 0, 0, 0)
_W3CDTF = "2026-01-01T00:00:00Z"
_CORE_DATE_RE = re.compile(rb"(<dcterms:(created|modified)\b[^>]*>)[^<]*(</dcterms:\2>)")


def write_text(path: Path, text: str) -> Path:
    """Write UTF-8 text with LF line endings on every OS (offsets stay stable)."""
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    with open(path, "w", encoding="utf-8", newline="\n") as f:
        f.write(text)
    return path


# --- PDF ---------------------------------------------------------------------------


def _pdf_string(line: str) -> str:
    if not line.isascii():
        raise ValueError(f"PDF lines must be ASCII (Helvetica, standard encoding): {line!r}")
    return "(" + line.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)") + ")"


def pdf_bytes(pages: Sequence[Sequence[str]], title: str = "") -> bytes:
    """Build a minimal PDF 1.4: one Helvetica text stream per page, one line per entry.

    Object layout: 1 catalog, 2 page tree, 3 font, then a (page, content)
    pair per page, then the info dictionary. No timestamps, so the output is
    byte-identical for identical input.
    """
    if not pages:
        raise ValueError("a PDF needs at least one page")
    n = len(pages)
    page_ids = [4 + 2 * i for i in range(n)]
    content_ids = [5 + 2 * i for i in range(n)]
    info_id = 4 + 2 * n

    objects: dict[int, bytes] = {
        1: b"<< /Type /Catalog /Pages 2 0 R >>",
        2: (
            f"<< /Type /Pages /Kids [{' '.join(f'{p} 0 R' for p in page_ids)}] /Count {n} >>"
        ).encode("ascii"),
        3: b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica /Encoding /WinAnsiEncoding >>",
    }
    for i, lines in enumerate(pages):
        ops = ["BT", "/F1 11 Tf", "14 TL", "72 740 Td"]
        ops += [f"{_pdf_string(line)} Tj T*" for line in lines]
        ops.append("ET")
        stream = "\n".join(ops).encode("ascii")
        objects[content_ids[i]] = (
            f"<< /Length {len(stream)} >>\nstream\n".encode("ascii") + stream + b"\nendstream"
        )
        objects[page_ids[i]] = (
            "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            f"/Resources << /Font << /F1 3 0 R >> >> /Contents {content_ids[i]} 0 R >>"
        ).encode("ascii")
    objects[info_id] = (
        f"<< /Title {_pdf_string(title)} /Author {_pdf_string(AUTHOR)} "
        "/Producer (RootMind demo corpus writer) >>"
    ).encode("ascii")

    out = bytearray(b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n")
    offsets: dict[int, int] = {}
    for oid in sorted(objects):
        offsets[oid] = len(out)
        out += f"{oid} 0 obj\n".encode("ascii") + objects[oid] + b"\nendobj\n"
    xref_at = len(out)
    size = max(objects) + 1
    out += f"xref\n0 {size}\n".encode("ascii")
    out += b"0000000000 65535 f \n"
    for oid in range(1, size):
        out += f"{offsets[oid]:010d} 00000 n \n".encode("ascii")
    out += (
        f"trailer\n<< /Size {size} /Root 1 0 R /Info {info_id} 0 R >>\n"
        f"startxref\n{xref_at}\n%%EOF\n"
    ).encode("ascii")
    return bytes(out)


def write_pdf(path: Path, pages: Sequence[Sequence[str]], title: str = "") -> Path:
    """Write a PDF with one page per entry of ``pages`` (each a list of ASCII lines)."""
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(pdf_bytes(pages, title))
    return path


# --- Office (zip) files --------------------------------------------------------------


def _pin_zip(path: Path) -> None:
    """Rewrite an OOXML zip with fixed member timestamps and fixed core-property dates."""
    with zipfile.ZipFile(path) as src:
        members = [(info.filename, src.read(info.filename)) for info in src.infolist()]
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", compression=zipfile.ZIP_DEFLATED) as dst:
        for name, data in members:
            if name == "docProps/core.xml":
                data = _CORE_DATE_RE.sub(rb"\g<1>" + _W3CDTF.encode() + rb"\g<3>", data)
            info = zipfile.ZipInfo(name, date_time=_ZIP_TIME)
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = 0o644 << 16
            dst.writestr(info, data)
    path.write_bytes(buf.getvalue())


def _set_core_properties(props: Any, title: str) -> None:
    """python-docx and python-pptx share this core-properties interface."""
    props.title = title
    props.author = AUTHOR
    props.last_modified_by = AUTHOR
    props.comments = FOOTER
    props.created = FIXED_TIME
    props.modified = FIXED_TIME
    props.revision = 1


def write_docx(
    path: Path,
    title: str,
    paragraphs: Sequence[str],
    table: Sequence[Sequence[str]] | None = None,
) -> Path:
    """Write a DOCX: a title, then paragraphs (``"## text"`` becomes a heading), then a table."""
    from docx import Document

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    _set_core_properties(doc.core_properties, title)
    doc.add_heading(title, level=0)
    for paragraph in paragraphs:
        if paragraph.startswith("## "):
            doc.add_heading(paragraph[3:], level=1)
        else:
            doc.add_paragraph(paragraph)
    if table:
        grid = doc.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, value in enumerate(row):
                grid.cell(r, c).text = value
    doc.save(path)
    _pin_zip(path)
    return path


def write_xlsx(path: Path, sheets: Mapping[str, Sequence[Sequence[Any]]], title: str = "") -> Path:
    """Write an XLSX with one worksheet per mapping entry (rows of cell values)."""
    from openpyxl import Workbook

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(title=name)
        for row in rows:
            ws.append(list(row))
    wb.properties.title = title
    wb.properties.creator = AUTHOR
    wb.properties.lastModifiedBy = AUTHOR
    wb.properties.created = FIXED_TIME
    wb.save(path)
    _pin_zip(path)
    return path


def write_pptx(path: Path, slides: Sequence[tuple[str, Sequence[str]]], title: str = "") -> Path:
    """Write a PPTX: slide 1 is a title slide, the rest are title + bullet slides.

    Each entry is ``(slide title, lines)``. For the first slide the lines form
    the subtitle.
    """
    from pptx import Presentation

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    prs = Presentation()
    _set_core_properties(prs.core_properties, title)
    for index, (heading, lines) in enumerate(slides):
        layout = prs.slide_layouts[0 if index == 0 else 1]
        slide = prs.slides.add_slide(layout)
        slide.shapes.title.text = heading
        body = slide.placeholders[1]
        body.text_frame.text = "\n".join(lines)
    prs.save(path)
    _pin_zip(path)
    return path
